from docx import Document
import os,pprint
os.chdir(r'C:\Intel\SGC\试剂')
document = Document('原辅料组实验准备信息指南 Fourth_edition.docx')
tables = document.tables
#print(tables)
print(len(tables))

a = 1
b = {}
for table in tables[2::2]:
	c = len(table.columns)
	print('这是第'+str(a)+'个表')
	for i in range(len(table.rows)):
		result = table.cell(i,c-2).text
		item = table.cell(i,c-5).text
		b[result] = item
	#print('\n\n')
	a += 1
pprint.pprint(b)